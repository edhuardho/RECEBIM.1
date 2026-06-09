import streamlit as st
import pandas as pd
from datetime import datetime
import openpyxl
import io
import re
import streamlit.components.v1 as components

st.set_page_config(page_title="Controle de Recebimento", layout="wide")

# 🖥️ BANCO DE DADOS GLOBAL COMPARTILHADO (A nível de Servidor)
if "bd_global_cargas" not in st.session_state.__class__.__dict__:
    st.session_state.__class__.bd_global_cargas = pd.DataFrame(columns=["ID_CARGA", "NOME_CARGA", "CONTEUDO_CSV", "STATUS"])
if "bd_global_itens" not in st.session_state.__class__.__dict__:
    st.session_state.__class__.bd_global_itens = pd.DataFrame(columns=["ID_CARGA", "CODIGO", "DESCRICAO", "FABRICACAO", "VALIDADE", "QUANTIDADE", "CONFERENTE"])
if "modelos_excel_memoria" not in st.session_state.__class__.__dict__:
    st.session_state.__class__.modelos_excel_memoria = {}

# ----------------- INTERFACE SHIFT (MENU LATERAL) -----------------
st.sidebar.title("🎮 Controle de Acesso")
perfil = st.sidebar.radio("Selecione o seu Perfil:", ["🖥️ Painel do Supervisor (PC)", "📱 Conferência na Doca (Celular)"])

# Botão master para resetar o servidor caso queira limpar o dia de trabalho
st.sidebar.markdown("---")
if st.sidebar.button("🧹 Limpar Todas as Cargas do Sistema"):
    st.session_state.__class__.bd_global_cargas = pd.DataFrame(columns=["ID_CARGA", "NOME_CARGA", "CONTEUDO_CSV", "STATUS"])
    st.session_state.__class__.bd_global_itens = pd.DataFrame(columns=["ID_CARGA", "CODIGO", "DESCRICAO", "FABRICACAO", "VALIDADE", "QUANTIDADE", "CONFERENTE"])
    st.session_state.__class__.modelos_excel_memoria = {}
    st.rerun()

# ----------------- 1. TELA DO SUPERVISOR (PC) -----------------
if perfil == "🖥️ Painel do Supervisor (PC)":
    st.title("🖥️ Painel de Controle - Supervisor")
    st.write("Gerencie as cargas, envie os arquivos padrão e finalize o processo para gerar o Excel.")
    
    tab1, tab2 = st.tabs(["🆕 Criar/Nomear Nova Carga", "🏁 Validar e Finalizar Recebimento"])
    
    with tab1:
        st.subheader("Configurar Carga de Entrada")
        nome_carga = st.text_input("Digite o Nome/Número da Carga (Ex: Carga Tirolez 08/06):").strip()
        
        arquivo_csv = st.file_uploader("Suba o arquivo CSV bruto do fornecedor", type=["csv", "txt"])
        arquivo_modelo = st.file_uploader("Suba o Modelo Excel Padrão (SHELF - PADRÃO.xlsx)", type=["xlsx"])
        
        if st.button("🚀 Liberar Carga para a Doca"):
            if nome_carga and arquivo_csv is not None and arquivo_modelo is not None:
                try:
                    conteudo = arquivo_csv.read().decode("utf-8", errors="ignore")
                    linhas_texto = conteudo.splitlines()
                    
                    linhas_produtos = []
                    for linha in linhas_texto:
                        if not linha.strip():
                            continue
                        linha_limpa = linha.replace('"', '').replace('-[-[-[-[-[-[-[-[', '').replace('\t', ' ').strip()
                        match = re.search(r'(\d{4,8})\s*-\s*([^,]+)', linha_limpa)
                        
                        if match:
                            codigo = match.group(1).strip().lstrip('0')
                            descricao = match.group(2).strip()
                            
                            # 🛑 FILTRO DE SEGURANÇA: Bloqueia linhas de clientes, motoristas, relatórios e totais
                            termos_bloqueados = ["TOTAL", "EMISSÃO", "PÁGINA", "RELATÓRIO", "CLIENTE", "MOTORISTA", "FORNECEDOR", "CONHECIMENTO"]
                            if not any(t in descricao.upper() or t in linha_limpa.upper() for t in termos_bloqueados):
                                linhas_produtos.append({"Codigo_Prod": codigo, "Descricao_Prod": descricao})
                    
                    if len(linhas_produtos) > 0:
                        df_produtos = pd.DataFrame(linhas_produtos).drop_duplicates(subset=['Codigo_Prod'])
                        id_carga = str(int(datetime.now().timestamp()))
                        
                        # Injeta no banco global compartilhado do servidor
                        nova_carga = pd.DataFrame([{"ID_CARGA": id_carga, "NOME_CARGA": nome_carga, "CONTEUDO_CSV": df_produtos.to_json(orient="records"), "STATUS": "EM CONFERENCIA"}])
                        st.session_state.__class__.bd_global_cargas = pd.concat([st.session_state.__class__.bd_global_cargas, nova_carga], ignore_index=True)
                        
                        # Guarda o binário do modelo Excel associado a ID
                        st.session_state.__class__.modelos_excel_memoria[id_carga] = arquivo_modelo.read()
                        
                        st.success(f"🎉 Carga '{nome_carga}' liberada com sucesso! {len(df_produtos)} produtos limpos e disponíveis para a Doca.")
                    else:
                        st.error("⚠️ Nenhum produto válido encontrado no CSV.")
                except Exception as e:
                    st.error(f"Erro ao processar arquivos: {e}")
            else:
                st.warning("Preencha o nome da carga e envie ambos os arquivos.")

    with tab2:
        st.subheader("Acompanhamento em Tempo Real")
        df_cargas = st.session_state.__class__.bd_global_cargas
        cargas_ativas = df_cargas[df_cargas["STATUS"] == "EM CONFERENCIA"] if not df_cargas.empty else pd.DataFrame()
        
        if cargas_ativas.empty:
            st.info("Nenhuma carga ativa no momento.")
        else:
            escolha_carga = st.selectbox("Selecione a carga para verificar/fechar:", cargas_ativas["NOME_CARGA"].unique())
            row_carga = cargas_ativas[cargas_ativas["NOME_CARGA"] == escolha_carga].iloc[0]
            id_sel = str(row_carga["ID_CARGA"])
            
            df_bipado = st.session_state.__class__.bd_global_itens[st.session_state.__class__.bd_global_itens["ID_CARGA"].astype(str) == id_sel]
            
            st.write("### Itens conferidos pela equipe via celular:")
            st.dataframe(df_bipado)
            
            if st.button("🏁 Fechar Conta e Gerar Planilha Oficial"):
                if df_bipado.empty:
                    st.error("Nenhum item foi conferido nessa carga ainda.")
                else:
                    try:
                        modelo_bytes = st.session_state.__class__.modelos_excel_memoria.get(id_sel)
                        wb = openpyxl.load_workbook(io.BytesIO(modelo_bytes))
                        aba_nome = "SHELF" if "SHELF" in wb.sheetnames else wb.sheetnames[0]
                        ws = wb[aba_nome]
                        
                        linha_inicio = 5
                        for index, row in df_bipado.reset_index().iterrows():
                            linha_atual = linha_inicio + index
                            ws[f"A{linha_atual}"] = int(row["CODIGO"]) if str(row["CODIGO"]).isdigit() else row["CODIGO"]
                            ws[f"B{linha_atual}"] = row["DESCRICAO"]
                            ws[f"C{linha_atual}"] = row["FABRICACAO"]
                            ws[f"D{linha_atual}"] = row["VALIDADE"]
                            ws[f"H{linha_atual}"] = row["QUANTIDADE"]
                            
                            ws[f"E{linha_atual}"] = f"=D{linha_atual}-TODAY()"
                            ws[f"F{linha_atual}"] = f"=D{linha_atual}-C{linha_atual}"
                            ws[f"G{linha_atual}"] = f"=E{linha_atual}/F{linha_atual}"
                        
                        buffer = io.BytesIO()
                        wb.save(buffer)
                        buffer.seek(0)
                        
                        st.session_state.__class__.bd_global_cargas.loc[st.session_state.__class__.bd_global_cargas["ID_CARGA"].astype(str) == id_sel, "STATUS"] = "FINALIZADO"
                        
                        st.success("🎉 Planilha gerada com sucesso!")
                        st.download_button(
                            label="📥 Baixar Planilha de Controle Pronta",
                            data=buffer,
                            file_name=f"CONTROLE_SHELF_{escolha_carga.replace(' ', '_')}.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                        )
                    except Exception as e:
                        st.error(f"Erro ao gerar arquivo: {e}")

# ----------------- 2. TELA DO CONFERENTE (CELULAR) -----------------
else:
    st.title("📱 Conferência de Entrada - Doca")
    
    if st.button("🔄 Atualizar Lista de Cargas"):
        st.rerun()
        
    df_cargas = st.session_state.__class__.bd_global_cargas
    cargas_disponiveis = df_cargas[df_cargas["STATUS"] == "EM CONFERENCIA"] if not df_cargas.empty else pd.DataFrame()
    
    if cargas_disponiveis.empty:
        st.success("✅ Nenhuma carga pendente de conferência na doca!")
    else:
        carga_selecionada = st.selectbox("Selecione a Carga para conferir:", cargas_disponiveis["NOME_CARGA"].unique())
        row_c = cargas_disponiveis[cargas_disponiveis["NOME_CARGA"] == carga_selecionada].iloc[0]
        id_carga_ativa = str(row_c["ID_CARGA"])
        
        # Carrega e limpa os códigos do CSV original do fornecedor
        df_produtos_carga = pd.read_json(io.StringIO(row_c["CONTEUDO_CSV"]))
        df_produtos_carga["Codigo_Prod"] = df_produtos_carga["Codigo_Prod"].astype(str)
        
        # Busca o que já foi bipado para esta carga
        meus_itens_carga = st.session_state.__class__.bd_global_itens[st.session_state.__class__.bd_global_itens["ID_CARGA"] == id_carga_ativa]
        
        # 📊 CÁLCULO E BARRA DE PROGRESSO DA CARGA
        total_produtos = len(df_produtos_carga)
        codigos_conferidos = meus_itens_carga["CODIGO"].astype(str).unique() if not meus_itens_carga.empty else []
        produtos_conferidos = len(codigos_conferidos)
        produtos_faltantes = max(0, total_produtos - produtos_conferidos)
        
        percentual = (produtos_conferidos / total_produtos) if total_produtos > 0 else 0.0
        
        st.markdown("### 📈 Progresso da Conferência desta Carga")
        st.progress(percentual)
        st.caption(f"✅ **{produtos_conferidos}** conferidos | ⏳ **{produtos_faltantes}** restantes de um total de **{total_produtos}** itens.")
        
        # 🔍 VISUALIZADOR DE ITENS PENDENTES (O QUE FALTA CONFERIR)
        with st.expander("🔍 Visualizar Itens que FALTAM Conferir nesta Carga", expanded=False):
            df_faltantes = df_produtos_carga[~df_produtos_carga["Codigo_Prod"].isin(codigos_conferidos)]
            if df_faltantes.empty:
                st.success("🎉 Todos os itens desta carga já foram enviados!")
            else:
                st.dataframe(
                    df_faltantes.rename(columns={"Codigo_Prod": "Código", "Descricao_Prod": "Descrição do Produto"}),
                    use_container_width=True,
                    hide_index=True
                )
        
        st.markdown("---")
        conferente = st.text_input("Nome do Conferente:")
        codigo_bipado = st.text_input("Digite ou Bipe o Código do Produto:", autocomplete="off")
        
        if codigo_bipado:
            codigo_limpo = codigo_bipado.strip().lstrip('0')
            item = df_produtos_carga[df_produtos_carga["Codigo_Prod"] == codigo_limpo]
            
            if not item.empty:
                desc_item = item.iloc[0]["Descricao_Prod"]
                st.info(f"📦 **Item:** {desc_item}")
                
                tipo_unidade = st.radio("Tipo de Medida do Item:", ["Unidade (Un)", "Peso Variável (Kg)"], horizontal=True)
                
                with st.form(key="form_celular", clear_on_submit=True):
                    st.write("📅 **Preenchimento de Datas (Apenas números - DD/MM/AAAA)**")
                    
                    col1, col2 = st.columns(2)
                    with col1:
                        f_fab_txt = st.text_input("Data de Fabricação:", max_chars=10, placeholder="Ex: 08/06/2026")
                    with col2:
                        f_ven_txt = st.text_input("Data de Vencimento:", max_chars=10, placeholder="Ex: 08/09/2026")
                    
                    if tipo_unidade == "Peso Variável (Kg)":
                        st.write("⚖️ **Informações de Peso Variável**")
                        col_p1, col_p2 = st.columns(2)
                        with col_p1:
                            peso_medio = st.number_input("Peso Médio da Peça (Kg):", min_value=0.001, step=0.001, format="%.3f")
                        with col_p2:
                            qtd_pecas = st.number_input("Quantidade Total de Peças:", min_value=1, step=1)
                        
                        f_qtd = round(peso_medio * qtd_pecas, 3)
                        st.warning(f"Calculado automaticamente peso total de: **{f_qtd} Kg**")
                    else:
                        f_qtd = st.number_input("Quantidade de Volumes/Unidades:", min_value=1, step=1)
                        peso_medio = 0.0
                        qtd_pecas = 0
                        
                    btn_enviar = st.form_submit_button("🚀 Enviar Lote para o Supervisor")
                    
                    components.html(
                        """
                        <script>
                        var inputs = window.parent.document.querySelectorAll('input[type="text"]');
                        inputs.forEach(function(input) {
                            if(input.placeholder && input.placeholder.includes("Ex: 08/06")) {
                                input.setAttribute('inputmode', 'numeric');
                                input.addEventListener('input', function(e) {
                                    var x = e.target.value.replace(/\D/g, '').match(/(\d{0,2})(\d{0,2})(\d{0,4})/);
                                    e.target.value = !x[2] ? x[1] : x[1] + '/' + x[2] + (x[3] ? '/' + x[3] : '');
                                });
                            }
                        });
                        </script>
                        """,
                        height=0,
                    )
                    
                    if btn_enviar:
                        if not conferente:
                            st.error("⚠️ Preencha seu nome antes de enviar.")
                        elif len(f_fab_txt) < 10 or len(f_ven_txt) < 10:
                            st.error("⚠️ Digite as datas completas com dia, mês e ano (Ex: 08/06/2026).")
                        else:
                            try:
                                datetime.strptime(f_fab_txt, '%d/%m/%Y')
                                datetime.strptime(f_ven_txt, '%d/%m/%Y')
                                
                                desc_final = f"{desc_item} ({qtd_pecas} pçs x {peso_medio}kg)" if tipo_unidade == "Peso Variável (Kg)" else desc_item
                                
                                novo_lote = pd.DataFrame([{
                                    "ID_CARGA": id_carga_ativa,
                                    "CODIGO": codigo_limpo,
                                    "DESCRICAO": desc_final,
                                    "FABRICACAO": f_fab_txt,
                                    "VALIDADE": f_ven_txt,
                                    "QUANTIDADE": f_qtd,
                                    "CONFERENTE": conferente
                                }])
                                st.session_state.__class__.bd_global_itens = pd.concat([st.session_state.__class__.bd_global_itens, novo_lote], ignore_index=True)
                                st.success("✔️ Lote enviado com sucesso para o painel do supervisor!")
                                st.rerun()
                            except ValueError:
                                st.error("❌ Data inválida! Verifique os dias e meses digitados (Ex: não existe mês 13 ou dia 32).")
            else:
                st.error("❌ Código de produto não encontrado nesta carga.")
        
        st.markdown("---")
        st.subheader("📋 Meus Itens Enviados (Histórico)")
        if not meus_itens_carga.empty:
            st.dataframe(
                meus_itens_carga[["CODIGO", "DESCRICAO", "FABRICACAO", "VALIDADE", "QUANTIDADE", "CONFERENTE"]].rename(
                    columns={
                        "CODIGO": "Código",
                        "DESCRICAO": "Produto",
                        "FABRICACAO": "Fabricação",
                        "VALIDADE": "Vencimento",
                        "QUANTIDADE": "Qtd/Peso",
                        "CONFERENTE": "Conferente"
                    }
                ),
                use_container_width=True,
                hide_index=True
            )
        else:
            st.info("Nenhum item enviado para esta carga ainda.")
